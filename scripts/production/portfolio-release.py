"""Root-only release maintenance; no new endpoint, users, or schema changes."""
import hashlib,io,json,re,secrets,sys
from datetime import datetime,timedelta
from pathlib import Path
from uuid import UUID
sys.path.insert(0,'/app')
import sqlalchemy as sa
from sqlalchemy.orm import Session
from app import cms,models as m
from app.admin import now,digest,audit,SESSION_COOKIE,CSRF_COOKIE
from app.db import get_engine

PATCH={'title':'AVALLOW','description':'Infraestrutura financeira para crédito, automação e integração entre instituições.','tags':['Fintech','Crédito','Infraestrutura'],'developmentStage':'Em desenvolvimento','cover':'/images/avallow-fintech.png'}
def hashed(value):return hashlib.sha256(json.dumps(value,default=str,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
def target(db,lock=False):
    q=sa.select(m.cms_entries).where(m.cms_entries.c.module=='projetos',m.cms_entries.c.source_key=='initial:projetos:onstar-scd',m.cms_entries.c.deleted_at.is_(None))
    return dict(db.execute(q.with_for_update() if lock else q).mappings().one())
def others(db,id):
    return hashed([dict(r) for r in db.execute(sa.select(m.cms_entries).where(m.cms_entries.c.id!=id).order_by(m.cms_entries.c.id)).mappings()])
def cast(key,value):
    if value is None:return None
    if key.endswith('_at'):return datetime.fromisoformat(value) if isinstance(value,str) else value
    if key=='editor_id':return UUID(value) if isinstance(value,str) else value
    return value
def execute(mode,release):
    key='portfolio_release:'+release
    with Session(get_engine()) as db:
        row=target(db,mode in {'apply','rollback'})
        if mode=='snapshot':
            assert db.scalar(sa.text('SELECT version_num FROM alembic_version'))=='0003','Unexpected schema; no migration will run'
            assert db.scalar(sa.select(m.admin_users.c.id).where(m.admin_users.c.active.is_(True)).limit(1)),'An active admin is required'
            assert row['scheduled_at'] is None,'Scheduled AVALLOW requires review'
            print(json.dumps({'target':row,'others_hash':others(db,row['id'])},default=str));return
        if mode=='apply':
            before=json.load(sys.stdin)
            assert hashed(before['target'])==hashed(row),'Concurrent AVALLOW change; aborting'
            assert before['others_hash']==others(db,row['id']),'CMS changed since snapshot; aborting'
            if all((row['published_data'] or {}).get(k)==v for k,v in PATCH.items()):print('{"changed":false}');return
            assert row['scheduled_at'] is None
            if row['published_data']:public=dict(row['published_data'])
            else:
                seeds=json.loads(Path('/app/app/data/cms_initial.json').read_text())
                public=next(x['data'] for x in seeds if x['source_key']==row['source_key']).copy()
            # Publish only the approved public fields; preserve all pending drafts.
            public=cms.validate('projetos',{**public,**PATCH},True)
            draft=cms.validate('projetos',{**row['data'],**PATCH})
            if '_legacy' in row['data']:draft['_legacy']=row['data']['_legacy']
            updated=db.execute(m.cms_entries.update().where(m.cms_entries.c.id==row['id']).values(data=draft,title=PATCH['title'],version=row['version']+1).returning(m.cms_entries)).mappings().one()
            cms.refs(db,row['id'],'projetos',draft,'draft');cms.revision(db,updated,'release_avallow',None)
            published=cms.publish_entry(db,updated,None,public)
            receipt={'before':hashed(row),'after':hashed(dict(published)),'reverted':False}
            db.execute(m.site_settings.insert().values(key=key,value=receipt));db.commit()
            print('{"changed":true}');return
        if mode=='rollback':
            before=json.load(sys.stdin);receipt=db.scalar(sa.select(m.site_settings.c.value).where(m.site_settings.c.key==key))
            if not receipt or receipt.get('reverted'):print('No release data to revert');return
            assert receipt['before']==hashed(before['target']) and receipt['after']==hashed(row),'Concurrent edits; automatic data rollback refused'
            saved=before['target'];keys=['data','title','slug','status','published_data','published_at','scheduled_data','scheduled_at','editor_id','deleted_at']
            values={k:cast(k,saved[k]) for k in keys};values['version']=row['version']+1
            reverted=db.execute(m.cms_entries.update().where(m.cms_entries.c.id==row['id']).values(**values).returning(m.cms_entries)).mappings().one()
            for scope,data in [('draft',saved['data']),('published',saved['published_data']),('scheduled',saved['scheduled_data'])]:
                if data:cms.refs(db,row['id'],'projetos',data,scope)
                else:
                    for table in (m.cms_media_refs,m.cms_links):db.execute(table.delete().where(table.c.entry_id==row['id'],table.c.scope==scope))
            cms.revision(db,reverted,'rollback_release',None)
            db.execute(m.site_settings.update().where(m.site_settings.c.key==key).values(value={**receipt,'reverted':True}));db.commit();print('AVALLOW reverted; history retained');return
        if mode=='verify':
            before=json.load(sys.stdin)
            assert before['others_hash']==others(db,row['id']),'Other CMS entries changed; inspect before concluding'
            assert all((row['published_data'] or {}).get(k)==v for k,v in PATCH.items())
            print('AVALLOW persisted; other CMS entries unchanged');return
        raise ValueError('Unknown mode')

def probe(base,connect=None):
    """Public HTTPS authenticated staging round trip, then cancellation/revocation.

    No password read, no admin account created, no article saved or published.
    A five-minute audited session is issued for an existing active administrator.
    """
    import httpx
    from docx import Document
    from PIL import Image
    from app.docx_import import cancel_stage
    from app.config import get_settings
    assert base.rstrip('/')==get_settings().site_url.rstrip('/')
    # Disposable tests route through their Docker proxy, keeping the real Origin.
    # Production always connects to the actual configured public HTTPS address.
    if connect:assert get_settings().app_env!='production' and base.startswith('http://127.0.0.1:') and connect=='http://portfolio-proxy'
    token=secrets.token_urlsafe(32);csrf=secrets.token_urlsafe(32);stage=None;session_id=None;admin=None
    with Session(get_engine()) as db:
        admin=db.scalar(sa.select(m.admin_users.c.id).where(m.admin_users.c.active.is_(True)).limit(1));assert admin
        session_id=db.scalar(m.admin_sessions.insert().values(admin_id=admin,token_hash=digest(token),csrf_hash=digest(csrf),expires_at=now()+timedelta(minutes=5)).returning(m.admin_sessions.c.id))
        audit(db,'deployment.docx_probe_session',admin);db.commit()
    try:
        doc=Document();doc.add_heading('Homologação temporária DOCX',0);doc.add_paragraph('Documento sintético: não publicar.',style='Subtitle');doc.add_paragraph('Por Homologação local')
        for i in range(1,11):
            doc.add_heading(str(i)+'. Seção de homologação',1);doc.add_paragraph('Parágrafo de validação de ordem e persistência temporária.')
            if i in {1,3,6,9}:
                picture=io.BytesIO();Image.new('RGB',(480,240),(i*20,80,160)).save(picture,'PNG');picture.seek(0);doc.add_picture(picture);doc.add_paragraph('Figura '+str(i)+' — legenda de teste',style='Caption')
        raw=io.BytesIO();doc.save(raw)
        with httpx.Client(base_url=connect or base,timeout=45,follow_redirects=False) as client:
            client.cookies.set(SESSION_COOKIE,token);client.cookies.set(CSRF_COOKIE,csrf)
            headers={'Origin':base,'X-CSRF-Token':csrf}
            response=client.post('/api/admin/docx/preview',headers=headers,files={'file':('homologacao.docx',raw.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
            assert response.status_code==201,'Public DOCX upload failed: HTTP '+str(response.status_code)
            stage=response.json()['token'];data=response.json()
            assert data['sections']==10 and len(data['images'])==4
            render=client.post('/api/admin/docx/staging/'+stage+'/render',headers=headers,json={'data':data['data']});assert render.status_code==200
            body=render.json()['data']['body'];assert body.count('<figure>')==4 and body.count('<h2>')==10
            for image in data['images']:
                asset=client.get(image['previewUrl']);assert asset.status_code==200 and hashlib.sha256(asset.content).hexdigest()==image['sha256']
            panel=client.get('/admin/artigos');assert panel.status_code==200 and 'Importar artigo do Word' in panel.text
            home=client.get('/');assert home.status_code==200 and 'AVALLOW' in home.text and PATCH['description'] in home.text
            asset=client.get(PATCH['cover']);assert asset.status_code==200 and hashlib.sha256(asset.content).hexdigest()=='a1436552052ea70e34bd4f4da68d0c803eca93473fdae0d5c9077807e046b55a'
            assert client.delete('/api/admin/docx/staging/'+stage,headers=headers).status_code==200;stage=None
        label='PUBLIC HTTPS PASS' if base.startswith('https://') else 'LOCAL HTTP PASS'
        print(label+': AVALLOW, admin panel, DOCX upload, 10 sections, 4 images, preview; no article created')
    finally:
        try:
            if stage:
                with Session(get_engine()) as db:cancel_stage(stage,user={'admin_id':admin,'id':session_id},db=db)
        finally:
            with Session(get_engine()) as db:
                db.execute(m.admin_sessions.update().where(m.admin_sessions.c.id==session_id).values(revoked_at=now()));audit(db,'deployment.docx_probe_revoked',admin);db.commit()

if __name__=='__main__':
    if sys.argv[1]=='probe':probe(sys.argv[2].rstrip('/'),sys.argv[3] if len(sys.argv)>3 else None)
    else:execute(sys.argv[1],sys.argv[2])
