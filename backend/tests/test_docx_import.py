import io
import os
import re
import zipfile
from pathlib import Path
from uuid import uuid4
import pytest
import sqlalchemy as sa
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
from fastapi.testclient import TestClient
from app import models as m,cms
from app.admin import hasher
from app.db import get_engine
from app.docx_parser import parse_docx,DocxError,MAX_UPLOAD

def image_bytes(color='blue'):
    out=io.BytesIO();Image.new('RGB',(480,240),color).save(out,format='PNG');return out.getvalue()

def document(images=0,title='Synthetic editorial document'):
    doc=Document();doc.core_properties.author='';doc.add_paragraph(title,style='Title');doc.add_paragraph('Synthetic subtitle',style='Subtitle');doc.add_paragraph('Por Synthetic Author')
    doc.add_paragraph('Data editorial: 21/10/2026');doc.add_paragraph('Categorias: Tecnologia, Dados');doc.add_paragraph('Tags: Teste, Engenharia')
    doc.add_paragraph('An introduction for the synthetic document, deliberately distinct from its title and containing useful context.')
    for n in range(max(1,images)):
        doc.add_heading(str(n+1)+'. Test section',level=1);doc.add_paragraph('Before image '+str(n))
        if images:doc.add_picture(io.BytesIO(image_bytes(['blue','red','green','yellow'][n%4])));doc.add_paragraph('Figura '+str(n+1)+' – Diagram '+str(n),style='Caption')
        doc.add_paragraph('After image '+str(n));p=doc.add_paragraph();p.add_run('Bold text').bold=True;p.add_run(' and italic').italic=True
    doc.add_paragraph('First bullet',style='List Bullet');doc.add_paragraph('Second bullet',style='List Bullet');doc.add_paragraph('Number one',style='List Number')
    doc.add_paragraph('A highlighted quotation',style='Intense Quote')
    table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Column';table.cell(0,1).text='Value';table.cell(1,0).text='Preserved';table.cell(1,1).text='42'
    p=doc.add_paragraph();link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to('https://example.org/source',RT.HYPERLINK,is_external=True));run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='Verified source';run.append(text);link.append(run);p._p.append(link)
    doc.sections[0].footer.paragraphs[0].text='Private repeated footer must not appear'
    out=io.BytesIO();doc.save(out);return out.getvalue()

def mutate(blob,name,transform):
    result=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(blob)) as source,zipfile.ZipFile(result,'w',zipfile.ZIP_DEFLATED) as target:
        for part in source.infolist():target.writestr(part.filename,transform(source.read(part)) if part.filename==name else source.read(part))
    return result.getvalue()

def test_valid_metadata_and_semantics():
    result=parse_docx(document(4),'article.docx');data=result.data
    assert data['title']=='Synthetic editorial document' and data['subtitle']=='Synthetic subtitle' and data['author']=='Synthetic Author'
    assert data['publishedDate']=='2026-10-21' and data['category']=='Tecnologia' and data['tags']==['Teste','Engenharia']
    assert len(result.images)==4 and result.sections==4
    assert all(x in data['body'] for x in ['<strong>Bold text</strong>','<em> and italic</em>','<ul>','<ol','<table>','<blockquote>','https://example.org/source','<figcaption>'])
    assert 'Synthetic editorial document' not in data['body'] and 'Synthetic subtitle' not in data['body'] and 'Por Synthetic Author' not in data['body']
    assert 'Private repeated footer' not in data['body'];assert len(data['description'])<=280 and len(data['metaDescription'])<=160
    assert data['readingMinutes']>=1 and not data['planned']

@pytest.mark.parametrize('count',[0,1,4])
def test_images_remain_interleaved(count):
    result=parse_docx(document(count),'article.docx');assert len(result.images)==count
    for n,asset in enumerate(result.images):
        html=result.data['body'];assert html.index('Before image '+str(n))<html.index('/api/media/'+asset['id'])<html.index('Figura '+str(n+1))<html.index('After image '+str(n))
    assert 'base64' not in result.data['body']

@pytest.mark.parametrize('blob,name',[(b'broken','test.docx'),(b'PK\x03\x04invalid','test.docx'),(document(),'test.docm'),(b'x'*(MAX_UPLOAD+1),'test.docx')])
def test_invalid_inputs(blob,name):
    with pytest.raises(DocxError):parse_docx(blob,name)

def test_xxe_macros_path_traversal_and_zip_bomb():
    blob=document();xml=mutate(blob,'word/document.xml',lambda b:b.replace(b'<w:document',b'<!DOCTYPE document [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><w:document',1))
    with pytest.raises(DocxError):parse_docx(xml,'test.docx')
    for name,raw in [('word/vbaProject.bin',b'macro'),('../escape',b'bad'),('word/media/bomb.txt',b'0'*(2*1024*1024))]:
        out=io.BytesIO(blob)
        with zipfile.ZipFile(out,'a',zipfile.ZIP_DEFLATED) as z:z.writestr(name,raw)
        with pytest.raises(DocxError):parse_docx(out.getvalue(),'test.docx')

def test_missing_metadata():
    doc=Document();doc.core_properties.author='';doc.add_paragraph('Untitled source');doc.add_paragraph('A complete first introductory paragraph with sufficient text for a concise summary.');out=io.BytesIO();doc.save(out)
    result=parse_docx(out.getvalue(),'test.docx');assert result.data['author']=='' and result.data['publishedDate']=='' and result.warnings
    result=parse_docx(out.getvalue(),'test.docx',{'author':'Configured Author','category':'Configured category','tags':['Configured tag']});assert result.data['author']=='Configured Author'

def test_reference_document_if_provided():
    source=os.environ.get('DOCX_REFERENCE')
    if not source:pytest.skip('Owner reference is private and never committed')
    result=parse_docx(Path(source).read_bytes(),'reference.docx')
    assert result.data['title']=='O SPLIT DO MAL'
    assert result.data['author']=='Carlos André Cardoso dos Santos'
    assert 'SplitPayment' in result.data['tags']
    assert result.data['subtitle']=='Quando o imposto chega antes do dinheiro necessário para manter sua empresa viva'
    assert result.sections==10 and len(result.images)==4
    with zipfile.ZipFile(source) as archive:
        originals={__import__('hashlib').sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith('word/media/')}
    assert {image['sha256'] for image in result.images}==originals
    assert len(re.findall('<h2>',result.data['body']))==11
    assert all(str(n)+'. ' in result.data['body'] for n in range(1,11))
    assert result.data['body'].count('<figcaption>')==3 and '<blockquote>' in result.data['body']
    assert result.data['body'].count('<a href="https://')>=3 and 'O SPLIT DO MAL' not in result.data['body']

def test_plain_source_urls_are_safe_clickable_and_not_nested():
    doc=Document();doc.add_heading('Source links',0)
    doc.add_paragraph('Source (https://example.org/report). Next https://example.org/other?x=1&y=2;')
    out=io.BytesIO();doc.save(out)
    body=parse_docx(out.getvalue(),'source.docx').data['body']
    assert 'href="https://example.org/report"' in body and '>https://example.org/report</a>).' in body
    assert 'href="https://example.org/other?x=1&amp;y=2"' in body

def test_footnote_text_and_reference_are_preserved():
    doc=Document();doc.add_heading('Notes',0);doc.add_paragraph('Subtitle',style='Subtitle')
    paragraph=doc.add_paragraph('Editorial text with a note.')
    run=OxmlElement('w:r');note=OxmlElement('w:footnoteReference');note.set(qn('w:id'),'1');run.append(note);paragraph._p.append(run)
    original=io.BytesIO();doc.save(original);result=io.BytesIO()
    with zipfile.ZipFile(original) as archive,zipfile.ZipFile(result,'w',zipfile.ZIP_DEFLATED) as output:
        for item in archive.infolist():output.writestr(item,archive.read(item))
        output.writestr('word/footnotes.xml','<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>Documented source note</w:t></w:r></w:p></w:footnote></w:footnotes>')
    body=parse_docx(result.getvalue(),'notes.docx').data['body']
    assert 'href="#docx-note-1"' in body and 'id="docx-note-1"' in body and 'Documented source note' in body

@pytest.fixture
def client(tmp_path,monkeypatch):
    from app.main import app
    monkeypatch.setattr(cms,'UPLOAD_ROOT',tmp_path)
    email='docx-'+uuid4().hex+'@example.com';password='synthetic-password-32-characters'
    with get_engine().begin() as db:db.execute(m.admin_users.insert().values(email=email,password_hash=hasher.hash(password)))
    c=TestClient(app);c.get('/api/admin/auth/csrf');headers={'Origin':'http://localhost:8080','X-CSRF-Token':c.cookies['portfolio_csrf']}
    assert c.post('/api/admin/auth/login',json={'email':email,'password':password},headers=headers).status_code==200
    headers['X-CSRF-Token']=c.cookies['portfolio_csrf'];yield c,headers,tmp_path;c.close()

def stage(client,blob=None):
    c,h,_=client;r=c.post('/api/admin/docx/preview',files={'file':('test.docx',blob or document(4),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},headers=h);assert r.status_code==201,r.text;return r.json()

def confirm(client,preview,**changes):
    c,h,_=client;return c.post('/api/admin/docx/staging/'+preview['token']+'/commit',json={'data':preview['data'],'reviewed':True,**changes},headers=h)

def test_docx_eight_mib_boundary_and_excess_rejection(client):
    base=document(4)
    def padded(size):
        out=io.BytesIO(base)
        with zipfile.ZipFile(out,'a',zipfile.ZIP_STORED) as archive:archive.writestr('customXml/size-test.dat',b'')
        padding=size-len(out.getvalue());out=io.BytesIO(base)
        with zipfile.ZipFile(out,'a',zipfile.ZIP_STORED) as archive:archive.writestr('customXml/size-test.dat',os.urandom(padding))
        assert len(out.getvalue())==size
        return out.getvalue()
    blob=padded(MAX_UPLOAD);preview=stage(client,blob);assert len(preview['images'])==4
    c,h,root=client;before=set((root/'.docx-staging').iterdir())
    response=c.post('/api/admin/docx/preview',files={'file':('large.docx',blob+b'X')},headers=h)
    assert response.status_code==413 and response.headers['content-type'].startswith('application/json')
    assert set((root/'.docx-staging').iterdir())==before

def test_reference_real_upload_if_provided(client):
    source=os.environ.get('DOCX_REFERENCE')
    if not source:pytest.skip('Private reference not mounted')
    preview=stage(client,Path(source).read_bytes())
    assert preview['data']['title']=='O SPLIT DO MAL' and preview['sections']==10 and len(preview['images'])==4

def test_staging_draft_idempotency_and_media(client):
    from app.main import app
    c,h,root=client;p=stage(client,document(4,'Draft '+uuid4().hex));assert not list(root.glob('*.png'))
    assert c.get(p['images'][0]['previewUrl']).status_code==200
    r=confirm(client,p);assert r.status_code==200,r.text;entry=r.json()['entry'];assert entry['status']=='draft' and entry['published_data'] is None
    assert len(list(root.glob('*.png')))==4 and confirm(client,p).json()['entry']['id']==entry['id']
    assert c.get('/api/media/'+p['images'][0]['id']).status_code==200
    with TestClient(app) as anonymous:assert anonymous.get('/api/media/'+p['images'][0]['id']).status_code==401
    assert len(c.get('/api/admin/docx/history/'+entry['id']).json())==1
    assert c.post('/api/admin/cms/artigos/'+entry['id']+'/publish',json={'version':entry['version']},headers=h).status_code==200
    with TestClient(app) as anonymous:assert anonymous.get('/api/media/'+p['images'][0]['id']).status_code==200

def test_permissions_csrf_and_review_required(client):
    from app.main import app
    c,h,_=client
    with TestClient(app) as anonymous:assert anonymous.post('/api/admin/docx/preview',files={'file':('x.docx',document())}).status_code==401
    assert c.post('/api/admin/docx/preview',files={'file':('x.docx',document())}).status_code==403
    p=stage(client,document(0,'Permission '+uuid4().hex));assert c.post('/api/admin/docx/staging/'+p['token']+'/commit',json={'data':p['data']},headers=h).status_code==422

def test_slug_collision_version_preserves_public_snapshot(client):
    c,h,_=client;name='Collision '+uuid4().hex;blob=document(1,name);p=stage(client,blob);r=confirm(client,p).json()['entry']
    published=c.post('/api/admin/cms/artigos/'+r['id']+'/publish',json={'version':r['version']},headers=h).json()
    second=stage(client,blob);second['data']['description']='Reviewed revision'
    assert confirm(client,second).json()['entry']['id']==r['id']
    changed=confirm(client,second,decision='version',confirmCollision=True,targetId=r['id'],version=published['version']);assert changed.status_code==200,changed.text
    revision=changed.json()['entry'];assert revision['status']=='draft' and revision['published_data']==published['published_data']
    third=stage(client,blob);third['data']['slug']=third['suggestedSlug']
    new=confirm(client,third,decision='new',confirmCollision=True);assert new.status_code==200 and new.json()['entry']['id']!=r['id']

def test_storage_failure_no_partial_rows_or_files(client,monkeypatch):
    from app import docx_import as service
    c,h,root=client;p=stage(client,document(4,'Failure '+uuid4().hex));original=service.os.link;calls=[]
    def fail_second(source,dest):
        calls.append(dest)
        if len(calls)==2:raise OSError('Synthetic disk failure')
        return original(source,dest)
    monkeypatch.setattr(service.os,'link',fail_second);assert confirm(client,p).status_code==503
    assert list(root.glob('*.png'))==[]
    with get_engine().connect() as db:assert not db.scalar(sa.select(m.cms_entries.c.id).where(m.cms_entries.c.slug==p['data']['slug']))
    monkeypatch.setattr(service.os,'link',original);assert confirm(client,p).status_code==200

def test_replace_and_cancel(client):
    c,h,root=client;p=stage(client,document(1,'Replacement '+uuid4().hex));image=p['images'][0]
    replaced=c.put(image['previewUrl'],files={'file':('replacement.png',image_bytes('black'))},headers=h);assert replaced.status_code==200
    assert replaced.json()['images'][0]['sha256']!=image['sha256']
    assert c.delete('/api/admin/docx/staging/'+p['token'],headers=h).status_code==200
    assert list(root.glob('*.png'))==[]

def test_external_images_and_unsafe_links_are_not_executed(monkeypatch):
    import urllib.request
    monkeypatch.setattr(urllib.request,'urlopen',lambda *a,**k:(_ for _ in ()).throw(AssertionError('No network allowed')))
    blob=document(1)
    def external(raw):
        from lxml import etree
        tree=etree.fromstring(raw)
        for r in tree:
            if r.get('Type','').endswith('/image'):r.set('Target','https://example.invalid/remote.png');r.set('TargetMode','External')
            if r.get('Type','').endswith('/hyperlink'):r.set('Target','javascript:alert(1)')
        return etree.tostring(tree)
    result=parse_docx(mutate(blob,'word/_rels/document.xml.rels',external),'test.docx')
    assert result.images==[] and 'javascript:' not in result.data['body'] and result.warnings

def test_image_limit(monkeypatch):
    monkeypatch.setattr('app.docx_parser.MAX_IMAGES',2)
    with pytest.raises(DocxError):parse_docx(document(4),'test.docx')

@pytest.mark.parametrize('lost_ack',[False,True])
def test_commit_failure_and_ambiguous_ack_are_reconciled(client,monkeypatch,lost_ack):
    from sqlalchemy.orm import Session
    from app import docx_import as service
    import hashlib
    c,h,root=client;p=stage(client,document(1,'Database failure '+uuid4().hex));stage_hash=hashlib.sha256(p['token'].encode()).hexdigest();key='docx_stage:'+stage_hash
    original=Session.commit
    def failed(session):
        receipt=session.scalar(sa.select(m.site_settings.c.id).where(m.site_settings.c.key==key))
        if receipt:
            if lost_ack:original(session)
            raise RuntimeError('Synthetic commit failure')
        return original(session)
    monkeypatch.setattr(Session,'commit',failed);assert confirm(client,p).status_code==503
    assert len(list(root.glob('*.png')))==(1 if lost_ack else 0)
    monkeypatch.setattr(Session,'commit',original);assert confirm(client,p).status_code==200
    manifest_path=root/'.docx-staging'/stage_hash/'manifest.json'
    import json
    manifest=json.loads(manifest_path.read_text());manifest['expires']=0;service.write_manifest(manifest_path.parent,manifest)
    with Session(get_engine()) as db:service.cleanup_expired(db)
    assert not manifest_path.parent.exists() and len(list(root.glob('*.png')))==1

def test_crash_orphan_cleanup_deletes_only_owned_hard_links(client):
    from sqlalchemy.orm import Session
    from app import docx_import as service
    import hashlib,json
    c,h,root=client;p=stage(client,document(1,'Crash '+uuid4().hex));directory=root/'.docx-staging'/hashlib.sha256(p['token'].encode()).hexdigest()
    manifest=json.loads((directory/'manifest.json').read_text());item=manifest['images'][0]
    service.os.link(directory/item['key'],root/item['key']);unrelated=root/'unrelated.png';unrelated.write_bytes(b'preserve')
    manifest['expires']=0;manifest['recover_until']=0;service.write_manifest(directory,manifest)
    with Session(get_engine()) as db:service.cleanup_expired(db)
    assert not (root/item['key']).exists() and unrelated.read_bytes()==b'preserve'

def test_stale_version_is_never_overwritten(client):
    c,h,_=client;blob=document(0,'Stale '+uuid4().hex);p=stage(client,blob);first=confirm(client,p).json()['entry'];again=stage(client,blob)
    saved=c.put('/api/admin/cms/artigos/'+first['id'],json={'version':first['version'],'data':{**first['data'],'description':'Later user edit'}},headers=h).json()
    r=confirm(client,again,decision='version',confirmCollision=True,targetId=first['id'],version=first['version']);assert r.status_code==409
    assert c.get('/api/admin/cms/artigos/'+first['id']).json()['data']['description']=='Later user edit'

def test_stage_is_session_bound_and_bad_upload_writes_nothing(client):
    from app.main import app
    c,h,root=client
    assert c.post('/api/admin/docx/preview',files={'file':('bad.docx',b'broken')},headers=h).status_code==422
    assert c.post('/api/admin/docx/preview',files={'file':('big.docx',b'0'*(MAX_UPLOAD+1))},headers=h).status_code==413
    p=stage(client,document(0,'Session '+uuid4().hex));email=c.get('/api/admin/auth/me').json()['email']
    with TestClient(app) as other:
        other.get('/api/admin/auth/csrf');headers={'Origin':'http://localhost:8080','X-CSRF-Token':other.cookies['portfolio_csrf']}
        assert other.post('/api/admin/auth/login',json={'email':email,'password':'synthetic-password-32-characters'},headers=headers).status_code==200
        assert other.get('/api/admin/docx/staging/'+p['token']).status_code==403


def test_commit_receipt_recovery_and_archived_article(client):
    c,h,root=client;p=stage(client,document(4,'Receipt '+uuid4().hex));url='/api/admin/docx/staging/'+p['token']
    assert c.get(url+'/result').json()['status']=='not_committed'
    result=confirm(client,p).json();row=result['entry'];assert result['entryId']==row['id']
    assert c.get(url+'/result').json()['entryId']==row['id']
    archived=c.post('/api/admin/cms/artigos/'+row['id']+'/archive',headers=h,json={'version':row['version']}).json()
    assert confirm(client,p).json()['entry']['status']=='archived'
    assert c.post(url+'/recover',headers=h,json={'entryId':row['id'],'version':row['version']}).status_code==409
    recovered=c.post(url+'/recover',headers=h,json={'entryId':row['id'],'version':archived['version']})
    assert recovered.status_code==200,recovered.text
    restored=recovered.json()['entry'];assert restored['id']==row['id'] and restored['status']=='draft'
    assert restored['data']==row['data'] and len(list(root.glob('*.png')))==4
    assert len(c.get('/api/admin/docx/history/'+row['id']).json())==1
    assert c.post(url+'/recover',json={'entryId':row['id'],'version':restored['version']}).status_code==403
    import shutil
    shutil.rmtree(root/'.docx-staging')
    assert c.get(url+'/result').json()['entryId']==row['id']
    archived_again=c.post('/api/admin/cms/artigos/'+row['id']+'/archive',headers=h,json={'version':restored['version']}).json()
    assert c.post(url+'/recover',headers=h,json={'entryId':row['id'],'version':archived_again['version']}).json()['entry']['status']=='draft'
    from app.main import app
    with TestClient(app) as anonymous:assert anonymous.get(url+'/result').status_code==401


def test_receipt_is_bound_to_original_session(client):
    c,h,_=client;p=stage(client,document(0,'Private receipt '+uuid4().hex));entry=confirm(client,p).json()['entry']
    with get_engine().connect() as db:email=db.scalar(sa.select(m.admin_users.c.email).where(m.admin_users.c.id==__import__('uuid').UUID(entry['editor_id'])))
    from app.main import app
    with TestClient(app) as other:
        other.get('/api/admin/auth/csrf');headers={'Origin':'http://localhost:8080','X-CSRF-Token':other.cookies['portfolio_csrf']}
        assert other.post('/api/admin/auth/login',headers=headers,json={'email':email,'password':'synthetic-password-32-characters'}).status_code==200
        assert other.get('/api/admin/docx/staging/'+p['token']+'/result').status_code==403


def test_expired_staging_survives_cleanup_for_explicit_resume(client):
    import hashlib,json,time
    from sqlalchemy.orm import Session
    from app import docx_import as service
    c,h,root=client;p=stage(client,document(4,'Expiration '+uuid4().hex));url='/api/admin/docx/staging/'+p['token']
    edited={**p['data'],'description':'Editorial review must survive expiration'}
    assert c.post(url+'/render',headers=h,json={'data':edited}).status_code==200
    directory=root/'.docx-staging'/hashlib.sha256(p['token'].encode()).hexdigest()
    manifest=json.loads((directory/'manifest.json').read_text());manifest['expires']=time.time()-1;service.write_manifest(directory,manifest)
    assert confirm(client,p).status_code==410
    with Session(get_engine()) as db:service.cleanup_expired(db)
    assert directory.exists(),'cleanup erased a recoverable editorial review'
    resumed=c.post(url+'/resume',headers=h);assert resumed.status_code==200,resumed.text
    assert resumed.json()['data']['description']==edited['description'] and len(resumed.json()['images'])==4
    p['data']=resumed.json()['data'];assert confirm(client,p).status_code==200


def test_repeat_commit_after_staging_removed_returns_saved_article(client):
    import hashlib,shutil
    c,h,root=client;p=stage(client,document(4,'Repeat after removal '+uuid4().hex));entry=confirm(client,p).json()['entry']
    shutil.rmtree(root/'.docx-staging'/hashlib.sha256(p['token'].encode()).hexdigest())
    repeated=confirm(client,p);assert repeated.status_code==200,repeated.text
    assert repeated.json()['entryId']==entry['id'] and len(list(root.glob('*.png')))==4


def test_resume_requires_session_csrf_and_obeys_hard_deadline(client):
    import hashlib,json,time
    from app import docx_import as service
    from app.main import app
    c,h,root=client;p=stage(client,document(1,'Bounded recovery '+uuid4().hex));url='/api/admin/docx/staging/'+p['token']
    assert c.post(url+'/resume').status_code==403
    with TestClient(app) as anonymous:assert anonymous.post(url+'/resume',headers=h).status_code==401
    email=c.get('/api/admin/auth/me').json()['email']
    with TestClient(app) as other:
        other.get('/api/admin/auth/csrf');headers={'Origin':'http://localhost:8080','X-CSRF-Token':other.cookies['portfolio_csrf']}
        assert other.post('/api/admin/auth/login',json={'email':email,'password':'synthetic-password-32-characters'},headers=headers).status_code==200
        headers['X-CSRF-Token']=other.cookies['portfolio_csrf']
        assert other.post(url+'/resume',headers=headers).status_code==403
    directory=root/'.docx-staging'/hashlib.sha256(p['token'].encode()).hexdigest()
    manifest=json.loads((directory/'manifest.json').read_text());manifest['expires']=time.time()-1;manifest['recover_until']=time.time()-1;service.write_manifest(directory,manifest)
    response=c.post(url+'/resume',headers=h);assert response.status_code==410
    assert response.json()['detail']['code']=='docx_staging_gone'
    assert c.get(p['images'][0]['previewUrl']).status_code==410


def test_missing_staged_image_has_recoverable_error_not_generic_failure(client):
    import hashlib,json
    c,h,root=client;p=stage(client,document(1,'Missing image '+uuid4().hex))
    directory=root/'.docx-staging'/hashlib.sha256(p['token'].encode()).hexdigest()
    manifest=json.loads((directory/'manifest.json').read_text());(directory/manifest['images'][0]['key']).unlink()
    result=confirm(client,p);assert result.status_code==404
    assert result.json()['detail']['code']=='docx_staging_missing'
    with get_engine().connect() as db:assert not db.scalar(sa.select(m.cms_entries.c.id).where(m.cms_entries.c.slug==p['data']['slug']))


def test_receipt_for_soft_deleted_article_requires_explicit_recovery(client):
    c,h,root=client;p=stage(client,document(4,'Deleted receipt '+uuid4().hex));first=confirm(client,p).json()['entry']
    assert c.delete('/api/admin/cms/artigos/'+first['id'],params={'version':first['version']},headers=h).status_code==200
    repeated=confirm(client,p);assert repeated.status_code==200,repeated.text
    row=repeated.json()['entry'];assert row['id']==first['id'] and row['deleted_at'] and row['status']=='archived'
    assert c.get('/api/admin/cms/artigos/'+first['id']).status_code==404
    restored=c.post('/api/admin/docx/staging/'+p['token']+'/recover',headers=h,json={'entryId':row['id'],'version':row['version']})
    assert restored.status_code==200,restored.text
    assert restored.json()['entry']['deleted_at'] is None and restored.json()['entry']['data']==first['data']
    assert c.get('/api/admin/cms/artigos/'+first['id']).json()['status']=='draft'
    assert len(list(root.glob('*.png')))==4 and len(c.get('/api/admin/docx/history/'+first['id']).json())==1
